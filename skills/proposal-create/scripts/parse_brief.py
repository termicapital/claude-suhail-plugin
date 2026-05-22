"""Parse a project brief into the canonical proposal-section JSON schema.

Reads a .docx, .md, or .txt file and splits content by heading patterns
(Arabic + English) into the 11 canonical sections. Missing sections are
emitted as null so the skill can render explicit placeholders rather than
hallucinated content.

Usage:
    python parse_brief.py "C:\\path\\to\\brief.docx"
"""
from __future__ import annotations

import json
import os
import re
import sys
from typing import Any

SECTIONS = [
    "cover",
    "about_client",
    "opportunity",
    "solution",
    "scope",
    "methodology",
    "timeline",
    "team",
    "investment",
    "next_steps",
    "contact",
]

# Heading aliases — match liberally. Arabic forms first since briefs are usually Arabic.
HEADINGS = {
    "cover":        [r"الغلاف", r"عنوان", r"cover", r"title"],
    "about_client": [r"نبذة", r"عن العميل", r"عن المنظمة", r"about", r"client"],
    "opportunity":  [r"الفرصة", r"المشكلة", r"التحدي", r"opportunity", r"problem", r"challenge"],
    "solution":     [r"الحل", r"الحل المقترح", r"المقترح", r"solution", r"approach"],
    "scope":        [r"النطاق", r"المخرجات", r"التسليمات", r"scope", r"deliverables"],
    "methodology":  [r"المنهجية", r"المراحل", r"الطريقة", r"methodology", r"phases", r"method"],
    "timeline":     [r"الجدول الزمني", r"الجدول", r"المدة", r"timeline", r"schedule"],
    "team":         [r"الفريق", r"team"],
    "investment":   [r"الاستثمار", r"السعر", r"التكلفة", r"الميزانية", r"investment", r"pricing", r"cost", r"budget"],
    "next_steps":   [r"الخطوات التالية", r"الخطوات", r"next steps", r"next"],
    "contact":      [r"التواصل", r"الاتصال", r"contact"],
}


def read_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        try:
            import docx  # python-docx
        except ImportError:
            print("python-docx not installed; run: pip install python-docx", file=sys.stderr)
            sys.exit(2)
        doc = docx.Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    if ext in {".md", ".txt"}:
        with open(path, "r", encoding="utf-8") as f:
            return f.read()
    print(f"unsupported brief extension: {ext}", file=sys.stderr)
    sys.exit(2)


def section_for_line(line: str) -> str | None:
    norm = line.strip().lower()
    if not norm:
        return None
    for section, patterns in HEADINGS.items():
        for p in patterns:
            if re.search(p, norm):
                return section
    return None


def is_heading(line: str) -> bool:
    s = line.strip()
    if not s:
        return False
    if s.startswith("#"):
        return True
    if len(s) < 60 and not s.endswith((".", "؟", "!", "؛")) and s == s.strip(":-—"):
        return True
    return False


def split_sections(text: str) -> dict[str, list[str]]:
    buckets: dict[str, list[str]] = {s: [] for s in SECTIONS}
    current: str | None = None
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            if current:
                buckets[current].append("")
            continue
        if is_heading(line):
            sec = section_for_line(line)
            if sec:
                current = sec
                continue
        if current:
            buckets[current].append(line)
    return buckets


def make_payload(buckets: dict[str, list[str]]) -> dict[str, Any]:
    def joined(sec: str) -> str | None:
        body = "\n".join(buckets[sec]).strip()
        return body or None

    def bullets(sec: str) -> list[str] | None:
        body = joined(sec)
        if not body:
            return None
        items = [re.sub(r"^[-•*\d\.\)\s]+", "", ln).strip() for ln in body.splitlines() if ln.strip()]
        items = [i for i in items if i]
        return items or None

    return {
        "cover":        {"title_ar": joined("cover"), "subtitle_ar": None, "date": None},
        "about_client": {"text_ar": joined("about_client"), "tags_ar": None},
        "opportunity":  {"text_ar": joined("opportunity")},
        "solution":     {"text_ar": joined("solution")},
        "scope":        {"items_ar": bullets("scope")},
        "methodology":  {"phases": None, "raw_ar": joined("methodology")},
        "timeline":     {"phases": None, "raw_ar": joined("timeline")},
        "team":         {"roles": None, "raw_ar": joined("team")},
        "investment":   {"rows": None, "total_ar": None, "raw_ar": joined("investment")},
        "next_steps":   {"items_ar": bullets("next_steps")},
        "contact":      {"name_ar": None, "email": None, "phone": None, "raw_ar": joined("contact")},
    }


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: parse_brief.py <path>", file=sys.stderr)
        return 2
    path = sys.argv[1]
    if not os.path.isfile(path):
        print(f"not a file: {path}", file=sys.stderr)
        return 2
    text = read_text(path)
    buckets = split_sections(text)
    print(json.dumps(make_payload(buckets), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
